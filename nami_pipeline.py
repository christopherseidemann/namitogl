#!/usr/bin/env python3
"""
NaMi Mitgliederlisten-Pipeline
Erstellt je eine XLSX-Datei pro Stufe mit Mitgliedern und Schnuppermitgliedern.

Verwendung:
    python nami_pipeline.py <nami_export.xlsx> [ausgabe_ordner]
"""

import sys
import os
import traceback
from datetime import datetime, timedelta

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Konfiguration ────────────────────────────────────────────────────────────
STUFEN = {
    "Wölfling":       "Wölfling",
    "Jungpfadfinder": "Jungpfadfinder",
    "Pfadfinder":     "Pfadfinder",
    "Rover":          "Rover",
}

SCHNUPPER_WARNUNG_WOCHEN = 6

# Farben
HEADER_BG      = "1F4E79"
HEADER_FG      = "FFFFFF"
ALT_ROW_BG     = "D6E4F0"
SCHNUPPER_ROT  = "C00000"  # Rote Schrift für überfällige Schnuppermitglieder

# ── Hilfsfunktionen ──────────────────────────────────────────────────────────

def datum_formatieren(wert):
    """Konvertiert Excel-Datumswert oder pandas-Timestamp in TT.MM.JJJJ."""
    if pd.isna(wert):
        return ""
    if isinstance(wert, (int, float)):
        try:
            d = datetime.fromordinal(datetime(1900, 1, 1).toordinal() + int(wert) - 2)
            return d.strftime("%d.%m.%Y")
        except Exception:
            return str(int(wert))
    if hasattr(wert, 'strftime'):
        return wert.strftime("%d.%m.%Y")
    return str(wert)


def zu_datetime(wert):
    """Gibt einen datetime zurück oder None."""
    if pd.isna(wert):
        return None
    if isinstance(wert, (int, float)):
        try:
            return datetime.fromordinal(datetime(1900, 1, 1).toordinal() + int(wert) - 2)
        except Exception:
            return None
    if hasattr(wert, 'to_pydatetime'):
        return wert.to_pydatetime()
    if isinstance(wert, datetime):
        return wert
    return None


def format_sheet(ws, stufen_name, schnupper_zeilen):
    """
    Formatiert das Worksheet.
    schnupper_zeilen: set von Excel-Zeilennummern (1-basiert) mit überfälliger Schnuppermitglied-Warnung.
    """
    thin = Side(border_style="thin", color="AAAAAA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Header
    header_fill = PatternFill("solid", start_color=HEADER_BG)
    for cell in ws[1]:
        cell.font = Font(bold=True, color=HEADER_FG, name="Arial", size=11)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    alt_fill = PatternFill("solid", start_color=ALT_ROW_BG)

    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        is_warnung = row_idx in schnupper_zeilen
        for cell in row:
            is_datum_spalte = cell.column == 3  # Spalte C = Anmeldedatum
            rot = is_warnung and is_datum_spalte
            cell.font = Font(
                name="Arial", size=10,
                color=SCHNUPPER_ROT if rot else "000000",
                bold=rot,
            )
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = border
            if row_idx % 2 == 0:
                cell.fill = alt_fill

    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max_len + 4

    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"
    ws.title = stufen_name


def stufe_xlsx_erstellen(df_stufe, ausgabe_pfad, stufen_name):
    heute = datetime.now()
    grenze = heute - timedelta(weeks=SCHNUPPER_WARNUNG_WOCHEN)

    zeilen = []
    schnupper_warnzeilen = set()  # Excel-Zeilennummern (Header = 1, Daten ab 2)

    for _, r in df_stufe.iterrows():
        ist_schnupper = r.get("Mitgliedstyp") == "Schnuppermitglied"
        eintritt = zu_datetime(r.get("Eintrittsdatum")) if ist_schnupper else None
        ueberfaellig = ist_schnupper and eintritt is not None and eintritt < grenze

        zeilen.append({
            "Vorname":       r.get("Vorname", "") or "",
            "Nachname":      r.get("Nachname", "") or "",
            "Anmeldedatum":  datum_formatieren(r.get("Eintrittsdatum")) if ist_schnupper else "",
            "_ueberfaellig": ueberfaellig,
        })

    df_out = pd.DataFrame(zeilen, columns=["Vorname", "Nachname", "Anmeldedatum", "_ueberfaellig"])
    if not df_out.empty:
        df_out = df_out.sort_values(["Nachname", "Vorname"]).reset_index(drop=True)

    # Warnzeilen bestimmen (nach Sortierung, Excel-Zeile = Index + 2)
    warnzeilen = set()
    for i, row in df_out.iterrows():
        if row["_ueberfaellig"]:
            warnzeilen.add(i + 2)

    df_out = df_out.drop(columns=["_ueberfaellig"])
    df_out.to_excel(ausgabe_pfad, index=False, sheet_name=stufen_name)

    wb = load_workbook(ausgabe_pfad)
    ws = wb.active
    format_sheet(ws, stufen_name, warnzeilen)
    wb.save(ausgabe_pfad)

    return len(df_out), len(warnzeilen)


def fehler_xlsx_erstellen(fehler_liste, ausgabe_pfad):
    if not fehler_liste:
        return 0
    df_fehler = pd.DataFrame(fehler_liste)
    df_fehler.to_excel(ausgabe_pfad, index=False, sheet_name="Fehler")
    wb = load_workbook(ausgabe_pfad)
    ws = wb.active
    format_sheet(ws, "Fehler", set())
    wb.save(ausgabe_pfad)
    return len(df_fehler)


# ── Haupt-Pipeline ───────────────────────────────────────────────────────────

def pipeline(eingabe_pfad, ausgabe_ordner="."):
    os.makedirs(ausgabe_ordner, exist_ok=True)

    print(f"[LESE] Lese: {eingabe_pfad}")
    engine = "xlrd" if eingabe_pfad.lower().endswith(".xls") else None
    df = pd.read_excel(eingabe_pfad, engine=engine)
    print(f"   → {len(df)} Zeilen geladen")

    # Mitglieder + Schnuppermitglieder, beide nur wenn Aktiv
    maske = (df["Status"] == "Aktiv") & (df["Mitgliedstyp"].isin(["Mitglied", "Schnuppermitglied"]))
    df_aktiv = df[maske].copy()
    print(f"   → {len(df_aktiv)} aktive Mitglieder/Schnuppermitglieder nach Filter")

    fehler = []
    bekannte_stufen = set(STUFEN.values())

    # Pflichtfelder prüfen
    for _, zeile in df_aktiv.iterrows():
        if pd.isna(zeile.get("Vorname")) or pd.isna(zeile.get("Nachname")):
            fehler.append({
                "Mitglieds-Nr.": zeile.get("Mitglieds-Nr.", ""),
                "Problem":       "Vorname oder Nachname fehlt",
                "Stufe":         zeile.get("Stufe/Abteilung", ""),
                "Mitgliedstyp":  zeile.get("Mitgliedstyp", ""),
            })

    # Pro Stufe eine Datei
    for anzeige_name, nami_name in STUFEN.items():
        df_stufe = df_aktiv[df_aktiv["Stufe/Abteilung"] == nami_name]
        dateiname = f"{anzeige_name}.xlsx"
        pfad = os.path.join(ausgabe_ordner, dateiname)
        anzahl, warnungen = stufe_xlsx_erstellen(df_stufe, pfad, anzeige_name)
        warn_str = f", davon {warnungen} Schnupper-Warnungen [WARNUNG]" if warnungen else ""
        print(f"   [OK] {anzeige_name}: {anzahl} Personen{warn_str} → {dateiname}")

    # Mitglieder ohne bekannte Stufe → Fehler
    unbekannte = df_aktiv[
        ~df_aktiv["Stufe/Abteilung"].isin(bekannte_stufen) |
        df_aktiv["Stufe/Abteilung"].isna()
    ]
    for _, zeile in unbekannte.iterrows():
        fehler.append({
            "Mitglieds-Nr.": zeile.get("Mitglieds-Nr.", ""),
            "Vorname":       zeile.get("Vorname", ""),
            "Nachname":      zeile.get("Nachname", ""),
            "Problem":       f"Unbekannte oder fehlende Stufe: '{zeile.get('Stufe/Abteilung', '')}'",
            "Mitgliedstyp":  zeile.get("Mitgliedstyp", ""),
        })

    fehler_pfad = os.path.join(ausgabe_ordner, "Fehler.xlsx")
    if fehler:
        anzahl_fehler = fehler_xlsx_erstellen(fehler, fehler_pfad)
        print(f"   [WARNUNG]  Fehler: {anzahl_fehler} Einträge → Fehler.xlsx")
    else:
        print("   [OK] Keine Fehler gefunden")

    print("\nPipeline abgeschlossen.")


# ── Entry Point ──────────────────────────────────────────────────────────────

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Verwendung: python nami_pipeline.py <nami_export.xlsx> [ausgabe_ordner]")
        sys.exit(1)

    eingabe = sys.argv[1]
    ausgabe = sys.argv[2] if len(sys.argv) > 2 else "ausgabe"

    try:
        pipeline(eingabe, ausgabe)
    except Exception as e:
        print(f"\n[FEHLER] Fehler: {e}")
        traceback.print_exc()
        sys.exit(1)
