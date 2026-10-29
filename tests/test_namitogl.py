"""Tests mit erfundenen Mitgliedern. Es werden keine echten Daten benutzt."""
import io
from datetime import datetime

import openpyxl
import pandas as pd
import pytest

import namitogl

HEUTE = datetime(2026, 10, 1)          # fest, damit die Schnupper-Warnung nicht vom Testtag abhängt
VOR_10_WOCHEN = datetime(2026, 7, 23)  # überfällig (Grenze: 6 Wochen)
VOR_2_WOCHEN = datetime(2026, 9, 17)   # noch nicht überfällig


def mitglied(nr, vorname, nachname, stufe, typ='Mitglied', status='Aktiv',
             eintritt=datetime(2020, 1, 15)):
    """Eine Zeile, wie sie im NaMi-Export vorkommt (gekürzt auf die nötigen Spalten)."""
    return {'Mitglieds-Nr.': nr, 'Vorname': vorname, 'Nachname': nachname,
            'Status': status, 'Mitgliedstyp': typ, 'Stufe/Abteilung': stufe,
            'Eintrittsdatum': eintritt, 'IBAN': 'DE30100000001111111111'}


def als_export(zeilen):
    """Schreibt die Zeilen als Excel-Datei und liest sie wieder ein, wie beim echten Export."""
    buf = io.BytesIO()
    pd.DataFrame(zeilen).to_excel(buf, index=False)
    buf.seek(0)
    return namitogl.lies_export(buf)


STANDARD = [
    mitglied(1, 'Anna', 'Müller', 'Wölfling'),
    mitglied(2, 'Ben', 'Schnupp', 'Wölfling', typ='Schnuppermitglied', eintritt=VOR_10_WOCHEN),
    mitglied(3, 'Dora', 'Frisch', 'Jungpfadfinder', typ='Schnuppermitglied', eintritt=VOR_2_WOCHEN),
    mitglied(4, 'Emil', 'Inaktiv', 'Rover', status='Inaktiv'),
    mitglied(5, 'Fritz', 'Ohnestufe', None),
]


def test_nur_aktive_mitglieder_und_schnuppermitglieder():
    ergebnis = namitogl.verarbeite(als_export(STANDARD), HEUTE)
    namen = [n for liste in ergebnis.listen.values() for n in liste['Vorname']]
    assert sorted(namen) == ['Anna', 'Ben', 'Dora']
    assert ergebnis.listen['Rover'].empty


def test_listen_enthalten_nur_namen_und_anmeldedatum():
    # Gruppenleitungen bekommen keine IBAN, Adresse oder Mitgliedsnummer
    ergebnis = namitogl.verarbeite(als_export(STANDARD), HEUTE)
    assert list(ergebnis.listen['Wölfling'].columns) == ['Vorname', 'Nachname', 'Anmeldedatum', 'Überfällig']


def test_anmeldedatum_nur_bei_schnuppermitgliedern():
    wl = namitogl.verarbeite(als_export(STANDARD), HEUTE).listen['Wölfling']
    assert wl.set_index('Vorname')['Anmeldedatum'].to_dict() == {'Anna': '', 'Ben': '23.07.2026'}


def test_ueberfaelliges_schnuppermitglied_wird_markiert():
    ergebnis = namitogl.verarbeite(als_export(STANDARD), HEUTE)
    assert ergebnis.listen['Wölfling'].set_index('Vorname')['Überfällig'].to_dict() == {'Anna': False, 'Ben': True}
    assert not ergebnis.listen['Jungpfadfinder']['Überfällig'].any()


def test_schnupper_warnung_auch_bei_datum_als_text():
    # In manchen Exporten steht das Datum als Text ('23.07.2026') statt als Datumszelle
    zeilen = [mitglied(6, 'Carl', 'Text', 'Wölfling', typ='Schnuppermitglied', eintritt='23.07.2026')]
    liste = namitogl.verarbeite(als_export(zeilen), HEUTE).listen['Wölfling']
    assert liste['Anmeldedatum'].tolist() == ['23.07.2026']
    assert liste['Überfällig'].tolist() == [True]


def test_sortierung_wie_im_telefonbuch():
    nachnamen = ['Zimmermann', 'Özdemir', 'de Vries', 'Müller', 'Ahrens']
    zeilen = [mitglied(i, 'X', n, 'Wölfling') for i, n in enumerate(nachnamen)]
    liste = namitogl.verarbeite(als_export(zeilen), HEUTE).listen['Wölfling']
    assert liste['Nachname'].tolist() == ['Ahrens', 'de Vries', 'Müller', 'Özdemir', 'Zimmermann']


def test_ueberfaellig_ist_in_excel_rot():
    liste = namitogl.verarbeite(als_export(STANDARD), HEUTE).listen['Wölfling']
    ws = openpyxl.load_workbook(io.BytesIO(namitogl.als_excel(liste, 'Wölfling'))).active
    zeilen = {r[0].value: r for r in ws.iter_rows(min_row=2)}
    assert zeilen['Ben'][2].font.color.rgb.endswith(namitogl.SCHNUPPER_ROT)
    assert not zeilen['Anna'][2].font.color.rgb.endswith(namitogl.SCHNUPPER_ROT)
    assert [c.value for c in ws[1]] == ['Vorname', 'Nachname', 'Anmeldedatum']   # Hilfsspalte nicht in der Datei


def test_mitglied_ohne_stufe_landet_in_der_fehlerliste():
    fehler = namitogl.verarbeite(als_export(STANDARD), HEUTE).fehler
    assert fehler['Nachname'].tolist() == ['Ohnestufe']


def test_fehlerliste_hat_einheitliche_spalten_und_lesbare_gruende():
    zeilen = [mitglied(7, None, 'Ohnevorname', 'Pfadfinder'),
              mitglied(8, 'Gustav', 'Ohnestufe', None),
              mitglied(9, 'Hanna', 'Leitung', 'Leiter')]
    fehler = namitogl.verarbeite(als_export(zeilen), HEUTE).fehler
    assert list(fehler.columns) == ['Problem', 'Mitglieds-Nr.', 'Vorname', 'Nachname',
                                    'Stufe/Abteilung', 'Mitgliedstyp']
    assert fehler['Problem'].tolist() == ['Vorname oder Nachname fehlt', 'Stufe fehlt',
                                          "Unbekannte Stufe: 'Leiter'"]


def test_kommandozeile_schreibt_eine_datei_pro_stufe(tmp_path):
    export = tmp_path / 'export.xlsx'
    pd.DataFrame(STANDARD).to_excel(export, index=False)
    namitogl.pipeline(str(export), str(tmp_path / 'ausgabe'))
    assert sorted(p.name for p in (tmp_path / 'ausgabe').iterdir()) == [
        'Fehler.xlsx', 'Jungpfadfinder.xlsx', 'Pfadfinder.xlsx', 'Rover.xlsx', 'Wölfling.xlsx']
