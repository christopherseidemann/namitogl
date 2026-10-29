#!/usr/bin/env python3
"""
namitogl: NaMi-Export -> eine Excel-Liste pro Stufe für die Gruppenleitungen.

Die Listen enthalten nur Vorname und Nachname, bei Schnuppermitgliedern
zusätzlich das Anmeldedatum. Überfällige Schnuppermitglieder sind rot markiert.

Verwendung:
    python namitogl.py <nami_export.xls> [ausgabe_ordner]
"""

import io
import os
import sys
import traceback
from dataclasses import dataclass
from datetime import datetime, timedelta

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ── Konfiguration ────────────────────────────────────────────────────────────
STUFEN = ["Wölfling", "Jungpfadfinder", "Pfadfinder", "Rover"]   # Werte der Spalte "Stufe/Abteilung"

SCHNUPPER_WARNUNG_WOCHEN = 6

# Farben
HEADER_BG      = "1F4E79"
HEADER_FG      = "FFFFFF"
ALT_ROW_BG     = "D6E4F0"
SCHNUPPER_ROT  = "C00000"  # Rote Schrift für überfällige Schnuppermitglieder


@dataclass
class Ergebnis:
    """Was bei der Verarbeitung herauskommt, noch ohne Dateien."""
    listen: dict          # Stufe -> DataFrame mit Vorname, Nachname, Anmeldedatum, Überfällig
    fehler: pd.DataFrame  # Einträge, die eine Gruppenleitung nicht zuordnen kann
    anzahl_gesamt: int = 0
    anzahl_aktiv:  int = 0


# ── Hilfsfunktionen ──────────────────────────────────────────────────────────

def datum_formatieren(wert):
    """Konvertiert Excel-Datumswert oder pandas-Timestamp in TT.MM.JJJJ."""
    if pd.isna(wert):
        return ""
    if isinstance(wert, (int, float)):
        try:
            d = datetime.fromordinal(datetime(1900, 1, 1).toordinal() + int(wert) - 2)
            return d.strftime("%d.%m.%Y")
        except Exception:
            return str(int(wert))
    if hasattr(wert, 'strftime'):
        return wert.strftime("%d.%m.%Y")
    return str(wert)


def zu_datetime(wert):
    """Gibt einen datetime zurück oder None."""
    if pd.isna(wert):
        return None
    if isinstance(wert, (int, float)):
        try:
            return datetime.fromordinal(datetime(1900, 1, 1).toordinal() + int(wert) - 2)
        except Exception:
            return None
    if hasattr(wert, 'to_pydatetime'):
        return wert.to_pydatetime()
    if isinstance(wert, datetime):
        return wert
    if isinstance(wert, str):
        # Datum als Text, z. B. '23.07.2026'. Festes Format, damit Tag und Monat nie vertauscht werden.
        try:
            return datetime.strptime(wert.strip(), "%d.%m.%Y")
        except ValueError:
            return None
    return None


def sortierschluessel(spalte):
    """Deutsche Sortierung (DIN 5007): Groß/klein egal, Ä wie A, Ö wie O, Ü wie U, ß wie ss."""
    return (spalte.fillna("").astype(str).str.casefold()     # casefold macht auch ß zu ss
            .str.replace("ä", "a").str.replace("ö", "o").str.replace("ü", "u"))


def lies_export(quelle):
    """Liest den NaMi-Export (.xls oder .xlsx). quelle = Dateipfad oder hochgeladene Datei."""
    return pd.read_excel(quelle)


# ── Verarbeitung ─────────────────────────────────────────────────────────────

def stufenliste(df_stufe, heute):
    """Liste für eine Stufe: Vorname, Nachname, Anmeldedatum, Überfällig."""
    grenze = heute - timedelta(weeks=SCHNUPPER_WARNUNG_WOCHEN)

    zeilen = []
    for _, r in df_stufe.iterrows():
        ist_schnupper = r.get("Mitgliedstyp") == "Schnuppermitglied"
        eintritt = zu_datetime(r.get("Eintrittsdatum")) if ist_schnupper else None
        ueberfaellig = ist_schnupper and eintritt is not None and eintritt < grenze

        zeilen.append({
            "Vorname":      r.get("Vorname", "") or "",
            "Nachname":     r.get("Nachname", "") or "",
            "Anmeldedatum": datum_formatieren(r.get("Eintrittsdatum")) if ist_schnupper else "",
            "Überfällig":   ueberfaellig,
        })

    liste = pd.DataFrame(zeilen, columns=["Vorname", "Nachname", "Anmeldedatum", "Überfällig"])
    if not liste.empty:
        liste = liste.sort_values(["Nachname", "Vorname"], key=sortierschluessel).reset_index(drop=True)
    return liste


def verarbeite(df, heute=None):
    """Teilt die aktiven Mitglieder auf die Stufen auf. Schreibt keine Dateien."""
    heute = heute or datetime.now()

    # Mitglieder + Schnuppermitglieder, beide nur wenn Aktiv
    maske = (df["Status"] == "Aktiv") & (df["Mitgliedstyp"].isin(["Mitglied", "Schnuppermitglied"]))
    df_aktiv = df[maske].copy()

    fehler = []

    def melde(zeile, problem):
        """Ein Eintrag der Fehlerliste, immer mit denselben Spalten."""
        fehler.append({"Problem": problem,
                       **{s: zeile.get(s) for s in ["Mitglieds-Nr.", "Vorname", "Nachname",
                                                    "Stufe/Abteilung", "Mitgliedstyp"]}})

    # Pflichtfelder prüfen
    for _, zeile in df_aktiv.iterrows():
        if pd.isna(zeile.get("Vorname")) or pd.isna(zeile.get("Nachname")):
            melde(zeile, "Vorname oder Nachname fehlt")

    # Pro Stufe eine Liste
    listen = {stufe: stufenliste(df_aktiv[df_aktiv["Stufe/Abteilung"] == stufe], heute)
              for stufe in STUFEN}

    # Mitglieder ohne bekannte Stufe → Fehler
    for _, zeile in df_aktiv[~df_aktiv["Stufe/Abteilung"].isin(STUFEN)].iterrows():
        stufe = zeile.get("Stufe/Abteilung")
        melde(zeile, "Stufe fehlt" if pd.isna(stufe) else f"Unbekannte Stufe: '{stufe}'")

    return Ergebnis(listen=listen, fehler=pd.DataFrame(fehler),
                    anzahl_gesamt=len(df), anzahl_aktiv=len(df_aktiv))


# ── Excel-Ausgabe ────────────────────────────────────────────────────────────

def als_excel(df, blattname):
    """Formatierte Excel-Datei als Bytes. Zeilen mit Überfällig=True: Anmeldedatum rot."""
    warn = df["Überfällig"].tolist() if "Überfällig" in df.columns else [False] * len(df)
    df = df.drop(columns=["Überfällig"], errors="ignore")

    wb = Workbook()
    ws = wb.active
    ws.title = blattname
    ws.append(list(df.columns))
    for werte in df.itertuples(index=False):
        ws.append([None if pd.isna(w) else w for w in werte])   # NaN -> leere Zelle

    thin = Side(border_style="thin", color="AAAAAA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Header
    header_fill = PatternFill("solid", start_color=HEADER_BG)
    for cell in ws[1]:
        cell.font = Font(bold=True, color=HEADER_FG, name="Arial", size=11)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    alt_fill = PatternFill("solid", start_color=ALT_ROW_BG)
    datum_spalte = list(df.columns).index("Anmeldedatum") + 1 if "Anmeldedatum" in df.columns else None

    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        is_warnung = warn[row_idx - 2]
        for cell in row:
            rot = is_warnung and cell.column == datum_spalte
            cell.font = Font(
                name="Arial", size=10,
                color=SCHNUPPER_ROT if rot else "000000",
                bold=rot,
            )
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = border
            if row_idx % 2 == 0:
                cell.fill = alt_fill

    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max_len + 4

    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def ausgabe_dateien(ergebnis):
    """Dateiname -> Inhalt (Bytes) für alle Listen und ggf. die Fehlerliste."""
    dateien = {f"{stufe}.xlsx": als_excel(liste, stufe) for stufe, liste in ergebnis.listen.items()}
    if not ergebnis.fehler.empty:
        dateien["Fehler.xlsx"] = als_excel(ergebnis.fehler, "Fehler")
    return dateien


# ── Kommandozeile ────────────────────────────────────────────────────────────

def pipeline(eingabe_pfad, ausgabe_ordner="."):
    os.makedirs(ausgabe_ordner, exist_ok=True)

    print(f"[LESE] Lese: {eingabe_pfad}")
    ergebnis = verarbeite(lies_export(eingabe_pfad))
    print(f"   → {ergebnis.anzahl_gesamt} Zeilen geladen")
    print(f"   → {ergebnis.anzahl_aktiv} aktive Mitglieder/Schnuppermitglieder nach Filter")

    for dateiname, inhalt in ausgabe_dateien(ergebnis).items():
        with open(os.path.join(ausgabe_ordner, dateiname), "wb") as f:
            f.write(inhalt)

    for stufe, liste in ergebnis.listen.items():
        warnungen = int(liste["Überfällig"].sum())
        warn_str = f", davon {warnungen} Schnupper-Warnungen [WARNUNG]" if warnungen else ""
        print(f"   [OK] {stufe}: {len(liste)} Personen{warn_str} → {stufe}.xlsx")

    if ergebnis.fehler.empty:
        print("   [OK] Keine Fehler gefunden")
    else:
        print(f"   [WARNUNG]  Fehler: {len(ergebnis.fehler)} Einträge → Fehler.xlsx")

    print("\nPipeline abgeschlossen.")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Verwendung: python namitogl.py <nami_export.xls> [ausgabe_ordner]")
        sys.exit(1)

    eingabe = sys.argv[1]
    ausgabe = sys.argv[2] if len(sys.argv) > 2 else "ausgabe"

    try:
        pipeline(eingabe, ausgabe)
    except Exception as e:
        print(f"\n[FEHLER] Fehler: {e}")
        traceback.print_exc()
        sys.exit(1)
